"""
Script : génère le fichier Excel des tarifs de livraison inter-quartiers Conakry
Usage  : python generate_tarifs_livraison.py
Sortie : Guimatrix_Tarifs_Livraison_Conakry.xlsx  (même dossier)
"""
from itertools import permutations
import os

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── Quartiers ─────────────────────────────────────────────────────────────────
QUARTIERS = [
    # Axe Carrefour
    "Madina", "Gbessiya", "Tanerie", "Sangoyah", "Tomboliya",
    "Lansanayah", "Km36", "Kabgelen", "Dubreka", "Gomboyah",
    "Friguiadi", "Somayah",
    # Axe Alignment
    "Dixcin", "Hamdalaye", "Cosa", "Enco 5", "Sonfonya",
    "Cimenterie", "Nationale",
    # Axe Corniche
    "Miniere", "Kipe", "Nongo", "Lambagni",
]

combos = list(permutations(QUARTIERS, 2))  # 506 paires dirigées

# ── Styles ────────────────────────────────────────────────────────────────────
HDR_FILL  = PatternFill("solid", fgColor="1a1a2e")
HDR_FONT  = Font(name="Arial", color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(name="Arial", size=10)
ALT_FILL  = PatternFill("solid", fgColor="EEF2FF")
YEL_FILL  = PatternFill("solid", fgColor="FFFDE7")   # cellules à remplir
CENTER    = Alignment(horizontal="center", vertical="center")
thin      = Side(style="thin", color="CCCCCC")
BORDER    = Border(left=thin, right=thin, top=thin, bottom=thin)

# ── Workbook ──────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()

# ─── Feuille 1 : Tarifs ───────────────────────────────────────────────────────
ws = wb.active
ws.title = "Tarifs"

headers = ["De (from_commune)", "Vers (to_commune)", "Tarif GNF ⭐ À REMPLIR", "Durée (heures)"]
widths  = [24, 24, 24, 18]

for col, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font      = HDR_FONT
    c.fill      = HDR_FILL
    c.alignment = CENTER
    c.border    = BORDER
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 30

for i, (src, dst) in enumerate(combos, 2):
    fill = ALT_FILL if i % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
    data = [src, dst, None, 1]
    for col, val in enumerate(data, 1):
        c = ws.cell(row=i, column=col, value=val)
        c.font      = BODY_FONT
        c.border    = BORDER
        c.alignment = CENTER
        c.fill      = YEL_FILL if col == 3 else fill

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:D{len(combos)+1}"

# ─── Feuille 2 : Légende ──────────────────────────────────────────────────────
ws2 = wb.create_sheet("Légende")
ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 55

infos = [
    ("Champ",                "Explication"),
    ("De (from_commune)",    "Quartier de départ du livreur / vendeur"),
    ("Vers (to_commune)",    "Quartier de destination / acheteur"),
    ("Tarif GNF",            "Prix de livraison en GNF (cases jaunes — à remplir)"),
    ("Durée (heures)",       "Délai estimé en heures (modifiable)"),
    ("",                     ""),
    ("Total combinaisons",   f"{len(combos)} paires (A→B et B→A séparés)"),
    ("Total quartiers",      str(len(QUARTIERS))),
    ("",                     ""),
    ("Étape suivante",       "Remplis les tarifs GNF, puis envoie le fichier à Claude"),
    ("",                     "Claude génèrera le script Django pour importer en base"),
]

hdr_font2 = Font(name="Arial", bold=True, size=11, color="FFFFFF")
for r, (k, v) in enumerate(infos, 1):
    ck = ws2.cell(row=r, column=1, value=k)
    cv = ws2.cell(row=r, column=2, value=v)
    if r == 1:
        ck.font = hdr_font2; cv.font = hdr_font2
        ck.fill = HDR_FILL;  cv.fill = HDR_FILL
    ck.font = Font(name="Arial", bold=(r==1), color=("FFFFFF" if r==1 else "000000"))
    cv.font = Font(name="Arial", color=("FFFFFF" if r==1 else "000000"))
    ck.border = BORDER; cv.border = BORDER

# ── Sauvegarder ───────────────────────────────────────────────────────────────
out = os.path.join(os.path.dirname(__file__), "Guimatrix_Tarifs_Livraison_Conakry.xlsx")
wb.save(out)
print(f"✅ Fichier créé : {out}")
print(f"   {len(combos)} combinaisons — {len(QUARTIERS)} quartiers")
print(f"   Remplis les cases jaunes (Tarif GNF), puis reviens pour le script d'import Django.")
